"""
Session 3 command. Builds the one-page monthly memo for leadership:

    python make_memo.py 2025-11

Uses only an APPROVED note (run approve.py first). Before building, it
recomputes the month's figures and checks every number in the approved note
still matches. If the data has changed since approval, it stops.

Saves in outputs/2025-11/:
  Zorora_MBR_2025-11.docx   editable Word file
  Zorora_MBR_2025-11.pdf    PDF (made through Word; if that fails, open the
                            Word file and use File > Save As > PDF)
"""

import json
import sys

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from approve import numbers_in
from src.drivers import check_events, drivers_for_month
from src.load import PROJECT_ROOT, load_data, load_settings
from src.locked_checks import run_locked_checks
from src.notes import build_slots
from src.report import MonthBuilder, fmt

GREY = RGBColor(0x59, 0x59, 0x59)

ROWS = [  # (metric, label, how the window is described)
    ("leads", "New leads", "month"),
    ("mqls", "MQLs (passed triage)", "month"),
    ("diagnostic_calls", "Diagnostic calls", "month"),
    ("paid_clients", "Paid clients", "month"),
    ("revenue", "Revenue", "month"),
    ("revenue_per_client", "Revenue per paid client", "window"),
    ("mql_rate", "MQL rate", "window"),
    ("close_rate", "Close rate", "window"),
    ("diag_to_close", "Diagnostic call to paid", "window"),
]
STATUS_TEXT = {"too_few": "too few to call", "launch_phase": "no comparison",
               "no_data": "no data"}


def cell_text(value, status, unit, rate):
    if status != "ok":
        return STATUS_TEXT.get(status, status)
    text = fmt(value, unit)
    if unit == "bdt":
        text += f" (CAD {round(value / rate):,})"
    return text


def main():
    if len(sys.argv) != 2:
        sys.exit("Usage: python make_memo.py YYYY-MM   (for example 2025-11)")
    month_text = sys.argv[1]
    folder = PROJECT_ROOT / "outputs" / month_text
    approved_file = folder / "notes_approved.json"
    if not approved_file.exists():
        sys.exit(f"No approved note for {month_text}. Run: python run_month.py {month_text}, "
                 f'then python approve.py {month_text} "Your Name"')
    approved = json.loads(approved_file.read_text(encoding="utf-8"))

    settings = load_settings()
    leads, revenue = load_data(settings)
    if (run_locked_checks(leads, revenue)["result"] == "FAIL").any():
        sys.exit("STOPPED: locked checks failed. Run python run_checks.py.")

    month = pd.Period(month_text, "M")
    report = MonthBuilder(leads, revenue, settings).report(month)
    drivers = drivers_for_month(report, month, check_events(leads, revenue), settings, leads)
    slots, _, _ = build_slots(report, month)

    # Re-check: every number in the approved note must still match today's figures
    allowed = set()
    for v in slots.values():
        allowed |= numbers_in(str(v))
    for d in drivers:
        allowed |= numbers_in(d["text"])
    stale = sorted({n for b in approved["bullets"] for n in numbers_in(b)} - allowed)
    if stale:
        sys.exit(f"STOPPED: the approved note has numbers that no longer match the data: "
                 f"{', '.join(stale)}. Run run_month.py and approve.py again.")
    print("Approved note re-checked: every number still matches the data.")

    rate = settings["bdt_per_cad"]
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Pt(54)
        section.left_margin = section.right_margin = Pt(60)

    doc.add_heading("Zorora Academy: Monthly Business Review", level=1)
    sub = doc.add_paragraph()
    run = sub.add_run(f"{month.strftime('%B %Y')}  |  Figures as of "
                      f"{month.end_time.strftime('%d %B %Y')}")
    run.font.color.rgb = GREY

    doc.add_heading("Key figures", level=2)
    t = report[report["segment"] == "all"].set_index("metric")
    table = doc.add_table(rows=1, cols=4)
    table.style = "Light Grid Accent 1"
    head = table.rows[0].cells
    for i, text in enumerate(["Figure", "This month", "Same month last year", "Last month"]):
        head[i].text = text
    for metric, label, kind in ROWS:
        r = t.loc[metric]
        name = label if kind == "month" else f"{label} ({r['window']})"
        cells = table.add_row().cells
        cells[0].text = name
        cells[1].text = cell_text(r["value"], r["status"], r["unit"], rate)
        cells[2].text = cell_text(r["last_year_value"], r["last_year_status"], r["unit"], rate)
        cells[3].text = cell_text(r["last_month_value"], r["last_month_status"], r["unit"], rate)

    doc.add_heading("Notes", level=2)
    for b in approved["bullets"]:
        doc.add_paragraph(b, style="List Bullet")
    sign = doc.add_paragraph()
    sign.add_run(f"Approved by {approved['approved_by']} on {approved['approved_on']}.").italic = True

    foot = doc.add_paragraph()
    foot.alignment = WD_ALIGN_PARAGRAPH.LEFT
    note = foot.add_run(
        "Data: reconstructed operating dataset calibrated to Zorora's real patterns, not a CRM "
        f"export. CAD shown for scale at 1 CAD = BDT {rate} ({settings['cad_rate_source']}, "
        f"{settings['cad_rate_date']}). Rates use rolling 3-month windows by the month leads "
        "arrived; close rate uses leads at least 30 days old; a rate with fewer than "
        f"{settings['min_denominator']} behind it is not shown. Comparison columns show what "
        "each earlier report showed for the same window. 2024 is the first comparison year.")
    note.font.size = Pt(8)
    note.font.color.rgb = GREY

    docx_path = folder / f"Zorora_MBR_{month_text}.docx"
    doc.save(docx_path)
    print(f"Saved outputs/{month_text}/{docx_path.name}")

    try:
        from docx2pdf import convert
        convert(str(docx_path), str(docx_path.with_suffix(".pdf")))
        print(f"Saved outputs/{month_text}/{docx_path.with_suffix('.pdf').name}")
    except Exception as e:  # Word missing, busy, or docx2pdf not installed
        print(f"PDF not made automatically ({type(e).__name__}). Open the Word file and use "
              "File > Save As > PDF.")


if __name__ == "__main__":
    main()
